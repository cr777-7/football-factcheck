#!/usr/bin/env python3
"""Render a human-verified football fact-check report as Markdown and optional XLSX."""
import argparse
import json
from pathlib import Path
from urllib.parse import urlparse
from parallel_check import read_claims

COLORS = {'错误': 'FCE4D6', '瑕疵': 'FFF2CC', '正确': 'E2F0D9', '证据不足': 'E7E6E6', '非事实判断': 'E7E6E6'}
FIELDS = [('id', 'ID'), ('text', '原断言'), ('context', '原文位置与语境'), ('status', '判定'), ('reason', '理由'), ('correction', '建议表述'), ('model_notes', '模型意见'), ('sources_text', '核验来源')]


def validate(report, original=None):
    if not isinstance(report, dict):
        raise ValueError('报告必须是对象')
    for key in ('title', 'checked_at', 'as_of', 'model_summary'):
        if not isinstance(report.get(key), str) or not report[key].strip():
            raise ValueError('缺少报告字段：' + key)
    if not isinstance(report.get('claims'), list) or not report['claims']:
        raise ValueError('报告需要claims数组')
    ids = set()
    for c in report['claims']:
        if not isinstance(c, dict):
            raise ValueError('断言必须是对象')
        for key in ('id', 'text', 'status', 'reason'):
            if not isinstance(c.get(key), str) or not c[key].strip():
                raise ValueError('缺少断言字段：' + key)
        if c['id'] in ids:
            raise ValueError('重复断言ID')
        ids.add(c['id'])
        if c['status'] not in COLORS:
            raise ValueError('未知判定')
        for key in ('correction', 'model_notes'):
            if not isinstance(c.get(key, ''), str):
                raise ValueError('无效文本字段')
        sources = c.get('sources', [])
        if not isinstance(sources, list):
            raise ValueError('sources必须是数组')
        verified = 0
        for s in sources:
            if not isinstance(s, dict):
                raise ValueError('来源必须是对象')
            for key in ('title', 'url', 'published_at', 'accessed_at', 'evidence'):
                if not isinstance(s.get(key), str) or not s[key].strip():
                    raise ValueError('来源缺少字段：' + key)
            url = urlparse(s['url'])
            if url.scheme not in ('https', 'http') or not url.hostname or url.username or url.password:
                raise ValueError('来源必须是有效网页URL')
            if s.get('verified') is True:
                verified += 1
        if c['status'] in ('错误', '瑕疵', '正确') and not verified:
            raise ValueError('事实判定必须附实际核验的网页证据')
    if original is not None:
        expected = {c['id']: c for c in original['claims']}
        if ids != set(expected):
            raise ValueError('报告与原始断言清单不一致：存在遗漏或新增ID')
        if report['as_of'] != original['as_of']:
            raise ValueError('报告适用时间与原断言清单不一致')
        for c in report['claims']:
            if c['text'] != expected[c['id']]['text']:
                raise ValueError('报告改写了原断言；修正应放入correction字段')
            c['context'] = expected[c['id']].get('context', '')
    return report


def source_text(c):
    return '\n\n'.join('%s\n%s\n发布：%s；访问：%s；已核验：%s\n%s' %
                        (s['title'], s['url'], s['published_at'], s['accessed_at'],
                         '是' if s.get('verified') is True else '否', s['evidence']) for s in c.get('sources', []))


def escaped(value):
    return str(value).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('|', '&#124;').replace('\n', '<br>')


def markdown(report):
    lines = ['# ' + escaped(report['title']), '', '核查时间：' + escaped(report['checked_at']) + '；适用时间：' + escaped(report['as_of']), '',
             '模型覆盖：' + escaped(report['model_summary']), '',
             '模型意见是线索，最终事实判定依据已核验的网页证据。', '',
             '| ' + ' | '.join(title for _, title in FIELDS) + ' |',
             '| ' + ' | '.join('---' for _ in FIELDS) + ' |']
    for c in report['claims']:
        row = dict(c, sources_text=source_text(c))
        lines.append('| ' + ' | '.join(escaped(row.get(key, '')) for key, _ in FIELDS) + ' |')
    lines.extend(['', '## 来源链接', ''])
    for c in report['claims']:
        for s in c.get('sources', []):
            title = escaped(s['title']).replace('[', '\\[').replace(']', '\\]')
            url = s['url'].replace(' ', '%20').replace('<', '%3C').replace('>', '%3E').replace('\n', '')
            lines.append('- ' + escaped(c['id']) + '：[' + title + '](<' + url + '>)')
    return '\n'.join(lines) + '\n'


def excel(report, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ValueError('Excel导出需要openpyxl，请安装requirements.txt')
    workbook = Workbook()
    summary = workbook.active
    summary.title = '报告说明'
    for row in [('标题', report['title']), ('核查时间', report['checked_at']), ('适用时间', report['as_of']),
                ('模型覆盖', report['model_summary']), ('说明', '模型审阅不是最终证据；来源核验由宿主助手完成。')]:
        if any(len(value) > 32767 for value in row):
            raise ValueError('Excel说明单元格超过32767字符，请缩短说明')
        summary.append(row)
    summary.column_dimensions['A'].width = 16
    summary.column_dimensions['B'].width = 100
    sheet = workbook.create_sheet('逐项核查')
    sheet.append([title for _, title in FIELDS])
    for c in report['claims']:
        row = dict(c, sources_text=source_text(c))
        values = [row.get(key, '') for key, _ in FIELDS]
        if any(len(value) > 32767 for value in values):
            raise ValueError('Excel单元格超过32767字符，请拆分报告；不会静默截断')
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.fill = PatternFill('solid', fgColor=COLORS[c['status']])
    for ws in workbook:
        for row in ws:
            for cell in row:
                # Force strings even when user content starts with =, +, -, or @.
                cell.data_type = 's'
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243746')
    for i, width in enumerate([12, 48, 40, 16, 55, 48, 40, 85], 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--claims', required=True, type=Path, help='最初的claims.json，用于检查遗漏及原文一致性')
    parser.add_argument('--markdown', type=Path)
    parser.add_argument('--xlsx', type=Path)
    args = parser.parse_args()
    if not args.markdown and not args.xlsx:
        parser.error('至少指定--markdown或--xlsx')
    try:
        paths = [x.resolve() for x in (args.input, args.claims, args.markdown, args.xlsx) if x]
        if len(paths) != len(set(paths)):
            raise ValueError('输入和输出路径必须各不相同')
        report = validate(json.loads(args.input.read_text(encoding='utf-8-sig')), read_claims(args.claims))
        if args.markdown:
            args.markdown.parent.mkdir(parents=True, exist_ok=True)
            args.markdown.write_text(markdown(report), encoding='utf-8')
        if args.xlsx:
            excel(report, args.xlsx)
    except (ValueError, OSError, TypeError) as error:
        parser.exit(2, '报告导出失败：' + (str(error) if type(error) is ValueError else type(error).__name__) + '\n')
    print('报告已生成。格式校验不代替来源核验。')


if __name__ == '__main__':
    main()
